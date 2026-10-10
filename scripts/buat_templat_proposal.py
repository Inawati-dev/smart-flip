"""Membuat empat templat Word Paket Rancangan Proposal (antrean #179).

Jalankan dari akar repo:  python scripts/buat_templat_proposal.py
Keluaran: public/templat/Templat-Bab-1-Pendahuluan.docx, ...-Bab-2-..., ...-Bab-3-...,
Templat-Rancangan-Proposal.docx. Judul bagian disalin dari canvas templat
(Templat-Bab1/2/3/Proposal) dan sama dengan src/lib/paketProposal.ts.
"""
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

KELUAR = Path(__file__).resolve().parent.parent / "public" / "templat"
FON = "Times New Roman"
ABU = RGBColor(0x80, 0x80, 0x80)


def set_fon(style, ukuran=12, tebal=None):
    style.font.name = FON
    style.font.size = Pt(ukuran)
    style.font.color.rgb = RGBColor(0, 0, 0)
    if tebal is not None:
        style.font.bold = tebal
    style.font.italic = False
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn("w:rFonts"))
    if rf is None:
        rf = OxmlElement("w:rFonts")
        rpr.insert(0, rf)
    for a in list(rf.attrib):  # buang fon tema bawaan (Calibri/Cambria)
        del rf.attrib[a]
    for a in ("ascii", "hAnsi", "eastAsia", "cs"):
        rf.set(qn("w:" + a), FON)


def baru():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Cm(21), Cm(29.7)
    s.left_margin = s.right_margin = s.top_margin = s.bottom_margin = Cm(2.54)
    set_fon(d.styles["Normal"])
    d.styles["Normal"].paragraph_format.line_spacing = 1.5
    d.styles["Normal"].paragraph_format.space_after = Pt(0)
    for nama, rata in (("Heading 1", WD_ALIGN_PARAGRAPH.CENTER), ("Heading 2", WD_ALIGN_PARAGRAPH.LEFT), ("Heading 3", WD_ALIGN_PARAGRAPH.LEFT)):
        st = d.styles[nama]
        set_fon(st, 12, True)
        pf = st.paragraph_format
        pf.alignment = rata
        pf.line_spacing = 1.5
        pf.space_before = Pt(12 if nama != "Heading 3" else 6)
        pf.space_after = Pt(0)
        pf.keep_with_next = True
    d.styles["Table Grid"].font.size = Pt(11)
    return d


def teks(p, isi, miring=False, abu=False, tebal=False, ukuran=None):
    r = p.add_run(isi)
    r.italic = miring
    r.bold = tebal
    if abu:
        r.font.color.rgb = ABU
    if ukuran:
        r.font.size = Pt(ukuran)
    return r


def petunjuk(d, isi):
    """Petunjuk pengisian: miring, abu-abu, dalam kurung siku. Dihapus mahasiswa."""
    p = d.add_paragraph()
    teks(p, "[" + isi + "]", miring=True, abu=True)
    return p


def bab(d, nomor, nama):
    p = d.add_paragraph(style="Heading 1")
    p.add_run(f"BAB {nomor}")
    p.add_run().add_break(WD_BREAK.LINE)
    p.add_run(nama.upper())


def bagian(d, isi, tingkat=2):
    return d.add_paragraph(isi, style=f"Heading {tingkat}")


def tabel(d, baris, lebar=None, kepala=False):
    """baris: list of list teks. Sel berawalan '[' diberi gaya petunjuk; selainnya tebal bila kolom 0 dan tidak kepala."""
    t = d.add_table(rows=len(baris), cols=len(baris[0]), style="Table Grid")
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(baris):
        for j, isi in enumerate(row):
            c = t.cell(i, j)
            if lebar:
                c.width = Cm(lebar[j])
            p = c.paragraphs[0]
            p.paragraph_format.line_spacing = 1.0
            p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(3)
            if isi.startswith("["):
                teks(p, isi, miring=True, abu=True)
            else:
                teks(p, isi, tebal=(kepala and i == 0) or (not kepala and j == 0))
    if kepala:  # baris judul diulang di halaman berikutnya
        trpr = t.rows[0]._tr.get_or_add_trPr()
        trpr.append(OxmlElement("w:tblHeader"))
    d.add_paragraph()
    return t


def kotak(d, isi):
    """Kotak bergaris putus-putus satu sel untuk tempat gambar."""
    t = d.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    tcpr = t.cell(0, 0)._tc.get_or_add_tcPr()
    b = OxmlElement("w:tcBorders")
    for sisi in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + sisi)
        e.set(qn("w:val"), "dashed")
        e.set(qn("w:sz"), "6")
        e.set(qn("w:color"), "808080")
        b.append(e)
    tcpr.append(b)
    p = t.cell(0, 0).paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = p.paragraph_format.space_after = Pt(18)
    teks(p, "[" + isi + "]", miring=True, abu=True)
    d.add_paragraph()


def data_diri(d, judul_proyek=True):
    for label in ("Nama", "NIM", "Kelas") + (("Judul Proyek",) if judul_proyek else ()):
        p = d.add_paragraph()
        p.paragraph_format.tab_stops.add_tab_stop(Cm(4))
        p.paragraph_format.line_spacing = 1.15
        teks(p, f"{label}\t: ", tebal=True)
        teks(p, "[ ... ]", miring=True, abu=True)
    d.add_paragraph()


def rujukan(d, isi):
    bagian(d, "RUJUKAN BAB INI")
    petunjuk(d, isi)


def lampiran(d, petunjuk_isi):
    p = bagian(d, "LAMPIRAN: HASIL AKTIVITAS MANDIRI", 1)
    p.paragraph_format.page_break_before = True
    petunjuk(d, petunjuk_isi)


def simpan(d, nama):
    KELUAR.mkdir(parents=True, exist_ok=True)
    d.save(KELUAR / nama)
    print("tulis", KELUAR / nama)


def pembuka(d, isi):
    petunjuk(d, isi + " Semua teks abu-abu miring dalam kurung siku adalah petunjuk: hapus seluruhnya sebelum berkas dikirim.")


FORMAT = "Format: Times New Roman 12, spasi 1,5, kertas A4. Simpan sebagai .docx lalu kirim lewat Mini Projek."


def bab1():
    d = baru()
    pembuka(d, "Templat Bab 1 Pendahuluan. Isi data diri, tulis tiap bagian di bawah judulnya, lalu tempel hasil Aktivitas Mandiri di Lampiran. " + FORMAT)
    data_diri(d)
    bab(d, 1, "Pendahuluan")
    bagian(d, "1.1 Latar Belakang Masalah")
    petunjuk(d, "Tiga paragraf atau lebih: konteks umum, lalu data dari observasi dan wawancara Anda, lalu mengapa masalah ini mendesak. Ambil dari AM-2 dan AM-3.")
    bagian(d, "1.2 Identifikasi Masalah")
    petunjuk(d, "Daftar masalah yang Anda temukan di lokasi, lalu masalah yang dipilih dan alasannya. Ambil dari AM-1 dan tabel skor di AM-2.")
    bagian(d, "1.3 Rumusan Masalah")
    petunjuk(d, "Satu sampai dua pertanyaan. Pola: Bagaimana mengembangkan ... untuk ... di ...? Bagaimana kelayakan ... menurut ... dan ...?")
    bagian(d, "1.4 Tujuan dan Manfaat")
    petunjuk(d, "Tujuan menjawab tiap rumusan masalah dengan kata kerja terukur. Manfaat ditulis dua: praktis dan teoretis. Ambil dari AM-3.")
    bagian(d, "1.5 Spesifikasi Produk yang Dikembangkan")
    tabel(d, [
        ["Nama Produk", "[ ... ]"],
        ["Bentuk", "[aplikasi, modul, alat, atau prosedur]"],
        ["Isi atau Fitur Utama", "[ ... ]"],
        ["Pengguna Sasaran", "[ ... ]"],
        ["Tempat Dipakai", "[ ... ]"],
        ["Batasan", "[yang tidak dikerjakan produk ini]"],
    ], lebar=[5, 11])
    rujukan(d, "Semua sumber yang dikutip di bab ini, format APA edisi ke-7. Nanti digabung ke Daftar Pustaka.")
    lampiran(d, "Halaman baru. Tempel hasil Aktivitas Mandiri apa adanya: AM-1 tabel masalah dan paradigma; AM-2 catatan observasi, wawancara, tabel skor tiga topik, studi pendahuluan; AM-3 draf pendahuluan. Tabel tulisan tangan boleh difoto lalu ditempel.")
    simpan(d, "Templat-Bab-1-Pendahuluan.docx")


def bab2():
    d = baru()
    pembuka(d, "Templat Bab 2 Landasan Teori. Isi data diri, tulis tiap bagian di bawah judulnya, lalu tempel hasil Aktivitas Mandiri di Lampiran. " + FORMAT)
    data_diri(d)
    bab(d, 2, "Landasan Teori")
    bagian(d, "2.1 Kajian Teori")
    bagian(d, "2.1.1 [Tema 1]", 3)
    petunjuk(d, "Satu paragraf sintesis atau lebih yang memadukan minimal dua sumber dengan parafrase. Paragraf dari AM-4 dipakai di sini.")
    bagian(d, "2.1.2 [Tema 2]", 3)
    petunjuk(d, "Sama seperti di atas, untuk tema kedua.")
    bagian(d, "2.1.3 [Tema 3, bila ada]", 3)
    petunjuk(d, "Hapus sub bab ini bila hanya ada dua tema.")
    bagian(d, "2.2 Penelitian yang Relevan")
    tabel(d, [["Peneliti (Tahun)", "Temuan Kunci", "Persamaan dengan Proyek Saya", "Perbedaan dengan Proyek Saya"]]
          + [["[ ... ]"] * 4 for _ in range(3)], lebar=[3.4, 4.2, 4.2, 4.2], kepala=True)
    petunjuk(d, "Minimal tiga penelitian. Tambah baris bila perlu.")
    bagian(d, "2.3 Kerangka Berpikir")
    tabel(d, [["Masalah di Lapangan", "Teori dan Penelitian yang Dipakai", "Produk yang Dikembangkan", "Hasil yang Diharapkan"]],
          lebar=[4, 4, 4, 4], kepala=True)
    petunjuk(d, "Isi tiap kotak dengan keadaan proyek Anda, beri panah dari kiri ke kanan, lalu jelaskan alurnya dalam satu paragraf.")
    rujukan(d, "Semua sumber yang dikutip di bab ini, format APA edisi ke-7. Tiap kutipan punya rujukan dan tiap rujukan dikutip.")
    lampiran(d, "Halaman baru. Tempel AM-4: daftar tema, lima sumber terkelompok, outline kajian pustaka.")
    simpan(d, "Templat-Bab-2-Landasan-Teori.docx")


def bab3():
    d = baru()
    pembuka(d, "Templat Bab 3 Model Pengembangan. Isi data diri, tulis tiap bagian di bawah judulnya, lalu tempel hasil Aktivitas Mandiri di Lampiran. " + FORMAT)
    data_diri(d)
    bab(d, 3, "Model Pengembangan")
    bagian(d, "3.1 Model Pengembangan dan Alasannya")
    petunjuk(d, "Sebut satu model yang dipakai dan alasannya dalam 3 sampai 5 kalimat, dikaitkan dengan jenis produk, tenggat, dan kebutuhan pengulangan. Ambil dari AM-5. Alasan memilih penelitian pengembangan ada di AM-1.")
    bagian(d, "3.2 Prosedur Pengembangan")
    petunjuk(d, "Ditulis sebagai narasi, satu sub bab per tahap model. Tiap paragraf menjawab: apa yang dikerjakan di proyek Anda, siapa yang terlibat, apa keluarannya, dan kapan.")
    bagian(d, "3.2.1 [Tahap 1 Model]", 3)
    petunjuk(d, "Satu paragraf. Contoh pembuka: Pada tahap ini saya ...")
    bagian(d, "3.2.2 [Tahap 2]", 3)
    petunjuk(d, "Satu paragraf.")
    bagian(d, "3.2.3 [Tahap Berikutnya]", 3)
    petunjuk(d, "Tambah sub bab sampai semua tahap model terbahas. Tabel tahap dari AM-5 cukup di lampiran.")
    bagian(d, "3.3 Validasi Ahli dan Uji Coba")
    petunjuk(d, "Siapa validator ahlinya, siapa sasaran uji coba pengguna, dan 3 sampai 5 kriteria yang dinilai validator. Ambil dari AM-5.")
    bagian(d, "3.4 Instrumen Pengumpulan Data")
    petunjuk(d, "Jenis instrumen (angket, pedoman wawancara, lembar observasi, atau gabungan) dan apa yang diukurnya. Instrumen lengkap minimal 8 butir diletakkan di lampiran. Ambil dari AM-6.")
    bagian(d, "3.5 Teknik Analisis Data")
    petunjuk(d, "Teknik yang dipakai (deskriptif, tematik, atau gabungan), siapa respondennya, kapan data dikumpulkan, dan cara menyimpulkannya. Ambil dari AM-6.")
    bagian(d, "3.6 Prototipe Produk")
    kotak(d, "Tempel gambar flowchart alur produk di sini, beri judul Gambar 3.1")
    petunjuk(d, "Tambahkan satu sampai tiga tangkapan layar prototipe dengan keterangan. Tautan prototipe (Figma atau sejenisnya) tidak ditulis di sini: isi di aplikasi saat mengirim.")
    rujukan(d, "Semua sumber yang dikutip di bab ini, format APA edisi ke-7.")
    lampiran(d, "Halaman baru. Tempel AM-5 (model, skema validasi, kriteria) dan AM-6 (instrumen lengkap, rencana analisis).")
    simpan(d, "Templat-Bab-3-Model-Pengembangan.docx")


def proposal():
    d = baru()
    pembuka(d, "Templat Rancangan Proposal Lengkap. Isi halaman judul, tempel Bab 1, 2, dan 3 yang sudah diperbaiki menurut umpan balik dosen, lalu susun Daftar Pustaka. Format: Times New Roman 12, spasi 1,5, kertas A4. Simpan sebagai satu berkas .docx lalu kirim lewat tombol Kirim Rancangan Proposal.")
    p = d.add_paragraph(style="Heading 1")
    p.add_run("RANCANGAN PROPOSAL")
    p.add_run().add_break(WD_BREAK.LINE)
    teks(p, "[Judul Proyek]", miring=True, abu=True).bold = False
    for label in ("Nama", "NIM", "Kelas", "Mata Kuliah", "Dosen Pengampu"):
        q = d.add_paragraph()
        q.paragraph_format.tab_stops.add_tab_stop(Cm(4.5))
        q.paragraph_format.line_spacing = 1.15
        teks(q, f"{label}\t: ", tebal=True)
        teks(q, "[ ... ]", miring=True, abu=True)
    d.add_paragraph()
    bagian(d, "Susunan Naskah")
    tabel(d, [
        ["Bab 1 Pendahuluan", "1.1 sampai 1.5, sudah diperbaiki menurut umpan balik"],
        ["Bab 2 Landasan Teori", "2.1 sampai 2.3, sudah diperbaiki"],
        ["Bab 3 Model Pengembangan", "3.1 sampai 3.6, sudah diperbaiki"],
        ["Daftar Pustaka", "Gabungan Rujukan Bab Ini dari ketiga bab"],
        ["Lampiran", "Instrumen lengkap. Hasil Aktivitas Mandiri tidak perlu dilampirkan lagi"],
    ], lebar=[6, 10])
    petunjuk(d, "Tempel ketiga bab berurutan sesudah halaman ini. Buang lampiran Aktivitas Mandiri dan bagian Rujukan Bab Ini dari tiap bab, isinya dipindah ke Daftar Pustaka. Hapus halaman petunjuk tempel di bawah setelah bab ditempel.")
    for n, nama in ((1, "Pendahuluan"), (2, "Landasan Teori"), (3, "Model Pengembangan")):
        q = d.add_paragraph(style="Heading 1")
        q.paragraph_format.page_break_before = True
        q.add_run(f"BAB {n} {nama.upper()}")
        petunjuk(d, f"Tempel Bab {n} di sini, sudah diperbaiki menurut umpan balik dosen.")
    q = bagian(d, "DAFTAR PUSTAKA", 1)
    q.paragraph_format.page_break_before = True
    petunjuk(d, "Halaman baru di akhir naskah. Cara menyusun: (1) kumpulkan Rujukan Bab Ini dari Bab 1, 2, 3; (2) buang yang kembar; (3) urutkan menurut abjad nama belakang penulis pertama; (4) tulis dengan format APA edisi ke-7, baris kedua dan seterusnya menjorok. Hanya sumber yang benar-benar dikutip di naskah.")
    pola = [
        ("Pola artikel jurnal: ", "Nama Belakang, Inisial., & Nama Belakang, Inisial. (Tahun). Judul artikel. ", "Nama Jurnal, volume", "(nomor), halaman. https://doi.org/..."),
        ("Pola buku: ", "Nama Belakang, Inisial. (Tahun). ", "Judul buku", " (edisi). Penerbit."),
        ("Pola laman atau dokumen lembaga: ", "Nama Lembaga. (Tahun). ", "Judul dokumen", ". Alamat laman"),
    ]
    for label, a, miring, b in pola:
        q = d.add_paragraph()
        q.paragraph_format.left_indent = Cm(1.27)
        q.paragraph_format.first_line_indent = Cm(-1.27)
        teks(q, label, miring=True, abu=True)
        teks(q, a)
        teks(q, miring, miring=True)
        teks(q, b)
    simpan(d, "Templat-Rancangan-Proposal.docx")


if __name__ == "__main__":
    bab1()
    bab2()
    bab3()
    proposal()
